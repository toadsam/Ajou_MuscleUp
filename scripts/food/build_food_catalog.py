"""식약처 K-FIND(가공식품 · 음식) + 농진청 국가표준식품성분 DB 에서 검색용 칸만 추려 foods.tsv.gz 를 만든다.
사용: python -I build_food_catalog.py <fooddata 폴더> <출력 파일>
칸: src(P 가공식품 · D 음식 · R 원재료) name maker basis(100g|100ml) kcal protein fat carb serving_g(1회 · 1인분) package_g(포장 전체)
"""
import gzip, re, sys, os
import openpyxl

src_dir, out_path = sys.argv[1], sys.argv[2]
rows = []


def num(v):
    if v is None:
        return None
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "tr", "Tr", "TR"):
        return None
    try:
        return round(float(s), 1)
    except ValueError:
        return None


def clean_maker(m):
    if not m:
        return ""
    m = str(m).strip()
    if m in ("해당없음", "-", "없음"):
        return ""
    m = re.sub(r"\(주\)|㈜|주식회사|\(유\)|유한회사|농업회사법인|영농조합법인", "", m)
    m = re.sub(r"\s+", " ", m).strip()
    # 같은 회사의 공장 · 지점은 회사 이름만(하림 제2공장 → 하림)
    m = re.sub(r"\s*제\s*\d+\s*공장$", "", m).strip()
    m = re.sub(r"\s+\S*(공장|지점|사업장)$", "", m).strip()
    return m


def grams(text):
    if not text:
        return None
    m = re.match(r"^\s*([\d.]+)\s*(g|ml|mL)\s*$", str(text))
    return round(float(m.group(1)), 1) if m else None


def kfind(path, src):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h else "" for h in next(it)]

    def col(*names):
        for n in names:
            for i, h in enumerate(header):
                if h.replace(" ", "") == n.replace(" ", ""):
                    return i
        raise KeyError(names)

    c_name, c_basis = col("식품명"), col("영양성분함량기준량")
    c_kcal, c_p, c_f, c_c = col("에너지(kcal)"), col("단백질(g)"), col("지방(g)"), col("탄수화물(g)")
    c_origin = col("식품기원명")

    def find(word):
        for i, h in enumerate(header):
            if word in h.replace(" ", ""):
                return i
        return None

    c_serv = find("참고량")
    c_weight = find("식품중량")
    c_mfr, c_dist, c_imp = find("제조사명"), find("유통업체명"), find("수입업체명")
    if c_mfr is None:
        c_mfr = find("업체명")
    print(src, "cols", c_serv, c_weight, c_mfr, c_dist, c_imp, flush=True)
    n = 0
    for r in it:
        name = (r[c_name] or "").strip() if r[c_name] else ""
        if not name:
            continue
        kcal = num(r[c_kcal])
        if kcal is None:
            continue
        g = lambda c: r[c] if c is not None else None
        maker = clean_maker(g(c_dist)) or clean_maker(g(c_mfr)) or clean_maker(g(c_imp))
        basis = str(r[c_basis] or "100g").strip().replace(" ", "")
        if basis not in ("100g", "100ml"):
            continue
        weight = grams(g(c_weight))  # 식품중량: 가공식품은 포장 전체, 음식은 1인분
        serving = grams(g(c_serv))   # 1회 섭취참고량
        if src == "D":
            serving, weight = serving or weight, None
        rows.append((src, name, maker, basis, kcal, num(r[c_p]), num(r[c_f]), num(r[c_c]),
                     serving, weight, str(r[c_origin] or "")))
        n += 1
    print(src, path, "rows", n, flush=True)


def rda(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    sheets = [ws for ws in wb.worksheets if ws.title.startswith("국가표준식품성분 Database")]
    ws = sorted(sheets, key=lambda w: w.title)[-1]
    print("RDA sheet", ws.title)
    n = 0
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i < 3 or not r[3]:
            continue
        kcal = num(r[5])
        if kcal is None:
            continue
        rows.append(("R", str(r[3]).strip(), "", "100g", kcal, num(r[7]), num(r[8]), num(r[10]), None, None, ""))
        n += 1
    print("R rows", n, flush=True)


rda(os.path.join(src_dir, "rda_raw.xlsx"))
kfind(os.path.join(src_dir, "kfind_02.xlsx"), "D")
kfind(os.path.join(src_dir, "kfind_01.xlsx"), "P")

# 음식 DB 는 같은 음식이 급식 · 외식 등으로 여러 줄 — 이름당 하나(급식이 아닌 것 우선)
out, seen = [], {}
for row in rows:
    src, name, maker = row[0], row[1], row[2]
    if src == "D":
        key = ("D", name.lower())
    else:
        # 공장 · 유통사만 다른 같은 제품은 하나로(이름과 영양값이 같으면)
        key = (src, re.sub(r"\s+", "", name.lower()), row[4], row[5], row[6], row[7])
    prefer = 0 if (src == "D" and "급식" in row[10]) else 1
    if key in seen:
        idx, p = seen[key]
        if prefer > p:
            out[idx] = row
            seen[key] = (idx, prefer)
        continue
    seen[key] = (len(out), prefer)
    out.append(row)


def fmt(v):
    if v is None:
        return ""
    return str(int(v)) if float(v).is_integer() else str(v)


with gzip.open(out_path, "wt", encoding="utf-8", newline="\n") as f:
    f.write("src\tname\tmaker\tbasis\tkcal\tprotein\tfat\tcarb\tserving_g\tpackage_g\n")
    for r in out:
        f.write("\t".join([r[0], r[1].replace("\t", " "), r[2].replace("\t", " "), r[3]] +
                          [fmt(x) for x in r[4:10]]) + "\n")
counts = {}
for r in out:
    counts[r[0]] = counts.get(r[0], 0) + 1
print("total", len(out), counts, "size", os.path.getsize(out_path))
for r in out:
    if "닭가슴" in r[1] and "하림" in r[2]:
        print("sample", r[:10])
        break
